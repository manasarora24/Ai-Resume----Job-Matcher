import re
from pathlib import Path

SKILLS = ["python", "javascript", "typescript", "react", "next.js", "node.js", "fastapi", "django", "flask", "sql", "postgresql", "mysql", "mongodb", "redis", "aws", "azure", "gcp", "docker", "kubernetes", "git", "machine learning", "deep learning", "pytorch", "tensorflow", "nlp", "scikit-learn", "pandas", "numpy", "data analysis", "tableau", "power bi", "html", "css", "tailwind", "java", "c++", "c#", "go", "rust", "agile", "communication", "leadership", "project management", "excel"]

def extract_text(data: bytes, suffix: str) -> str:
    if suffix == ".pdf":
        import io
        from pypdf import PdfReader
        text = "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
    elif suffix == ".docx":
        import io
        from docx import Document
        doc = Document(io.BytesIO(data))
        text = "\n".join([p.text for p in doc.paragraphs] + [cell.text for table in doc.tables for row in table.rows for cell in row.cells])
    else:
        raise ValueError("Only PDF and DOCX files are supported.")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < 40:
        raise ValueError("We could not extract readable text. Scanned PDFs need OCR before upload.")
    return text[:100_000]

def parse_resume(text: str) -> dict:
    lines = [x.strip(" •\t") for x in text.splitlines() if x.strip()]
    email = re.search(r"[\w.+-]+@[\w.-]+\.\w+", text)
    phone = re.search(r"(?:\+?\d[\d ().-]{7,}\d)", text)
    headings = {"summary": ["summary", "profile", "objective"], "experience": ["experience", "employment", "work history"], "education": ["education", "academics"], "projects": ["projects", "selected projects"], "certifications": ["certifications", "certificates"]}
    sections = {k: [] for k in headings}
    current = None
    for line in lines:
        key = next((k for k, names in headings.items() if line.lower().rstrip(":") in names), None)
        if key:
            current = key
        elif current and len(line) < 400:
            sections[current].append(line)
    low = text.lower()
    skills = [s for s in SKILLS if re.search(r"(?<![\w+#])" + re.escape(s.lower()) + r"(?![\w+#])", low)]
    return {"name": next((l for l in lines[:5] if "@" not in l and not re.search(r"\d{3}", l) and len(l.split()) <= 5), ""), "email": email.group(0) if email else "", "phone": phone.group(0) if phone else "", "skills": skills, **{k: "\n".join(v) for k,v in sections.items()}, "experience_years": None}
