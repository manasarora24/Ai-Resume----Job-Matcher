import os
import tempfile
from io import BytesIO
from pathlib import Path

_test_dir = Path(tempfile.mkdtemp(prefix="careeriq-tests-"))
os.environ["DATABASE_URL"] = "sqlite://"
os.environ["UPLOAD_DIR"] = str(Path(__file__).resolve().parents[1] / "uploads-test")
os.environ["JWT_SECRET"] = "unit-test-secret-not-for-production-123456789"

from fastapi.testclient import TestClient
from docx import Document
from app.main import app

client = TestClient(app)

def register(email):
    return client.post("/api/auth/register", json={"email": email, "password": "long-test-password", "name": "Test User"})

def test_registration_protection_and_account_scoping():
    a = register("alice@example.com")
    assert a.status_code == 201
    token_a = a.json()["token"]
    assert client.get("/api/jobs").status_code == 401
    job = client.post("/api/jobs", headers={"Authorization": f"Bearer {token_a}"}, json={
        "title": "Data Analyst", "company": "Example Co", "location": "Remote",
        "description": "Build Python and SQL data analysis workflows for the product team.",
    })
    assert job.status_code == 201
    document = Document()
    document.add_paragraph("Alex Example")
    document.add_paragraph("alex@example.com")
    document.add_paragraph("Experience")
    document.add_paragraph("Built Python and SQL data analysis workflows for product teams.")
    resume = client.post("/api/resumes", headers={"Authorization": f"Bearer {token_a}"}, files={
        "file": ("resume.docx", BytesIO(_docx_bytes(document)), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    })
    assert resume.status_code == 201
    assert "python" in resume.json()["parsed"]["skills"]
    profile = resume.json()["parsed"]
    profile["skills"] = ["python", "sql", "react"]
    edited = client.patch(f"/api/resumes/{resume.json()['id']}", headers={"Authorization": f"Bearer {token_a}"}, json={"parsed": profile})
    assert edited.status_code == 200
    result = client.post(f"/api/matches/{resume.json()['id']}/{job.json()['id']}", headers={"Authorization": f"Bearer {token_a}"})
    assert result.status_code == 200
    assert result.json()["matched_skills"] == ["data analysis", "python", "sql"]
    b = register("bob@example.com")
    token_b = b.json()["token"]
    headers_b = {"Authorization": f"Bearer {token_b}"}
    assert client.get("/api/jobs", headers=headers_b).json() == []
    assert client.post(f"/api/jobs/{job.json()['id']}/save", headers=headers_b).status_code == 404
    assert client.get(f"/api/resumes", headers=headers_b).json() == []

def _docx_bytes(document):
    data = BytesIO()
    document.save(data)
    return data.getvalue()

